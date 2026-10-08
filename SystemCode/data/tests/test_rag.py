"""Offline tests. A deterministic encoder/client exercises orchestration, not LLM quality."""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fleetcare_agent import FleetCareAssistant
from knowledge_rag import KnowledgeRetriever, build_index, document_units, tokens


class Tokenizer:
    def __init__(self):
        self.words = {}

    def num_special_tokens_to_add(self, pair=False):
        return 2

    def encode(self, text, **kwargs):
        result = []
        for word in text.split():
            self.words.setdefault(word, len(self.words) + 1)
            result.append(self.words[word])
        return result

    def decode(self, ids, **kwargs):
        reverse = {v: k for k, v in self.words.items()}
        return " ".join(reverse[i] for i in ids)


class Encoder:
    max_seq_length = 32

    def __init__(self):
        self.tokenizer = Tokenizer()

    def encode(self, texts, **kwargs):
        result = np.zeros((len(texts), 256), dtype=np.float32)
        for i, text in enumerate(texts):
            for word in tokens(text):
                index = int(hashlib.sha256(word.encode()).hexdigest()[:8], 16) % 256
                result[i, index] += 1
        return result


class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.encoder = Encoder()
        self.index = self.root / "index.npz"
        self.entries = []

    def add(self, name, content, kind="maintenance"):
        (self.root / name).write_text(content, encoding="utf-8")
        self.entries.append({"source": name, "title": name, "kind": kind,
                             "version": "test", "reviewed": True, "applicability": "test only"})
        (self.root / "catalog.json").write_text(json.dumps(self.entries), encoding="utf-8")

    def make(self):
        build_index(self.root, self.index, "test-encoder", encoder=self.encoder)
        return KnowledgeRetriever(self.root, self.index, encoder=self.encoder, min_similarity=0.4)

    def test_retrieval_and_scope_filter(self):
        self.add("brakes.md", "# Brakes\nBrake air leaks require inspection.")
        self.add("policy.md", "# Policy\nPriority inspection brake air leaks.", "project")
        result = self.make().search("brake air leaks", kind="maintenance")
        self.assertTrue(result["maintenance_evidence_found"])
        self.assertTrue(all(h["kind"] == "maintenance" for h in result["hits"]))
        self.assertEqual(result["hits"][0]["source"], "brakes.md")

    def test_unrelated_query_returns_no_evidence(self):
        self.add("brakes.md", "Brake air leaks")
        self.assertEqual(self.make().search("pineapple astronomy violin")["hits"], [])

    def test_modified_document_requires_rebuild(self):
        self.add("brakes.md", "Brake air leaks")
        retriever = self.make()
        (self.root / "brakes.md").write_text("Changed brakes", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "changed"):
            retriever.search("brakes")

    def test_changed_catalog_requires_rebuild(self):
        self.add("brakes.md", "Brake air leaks")
        retriever = self.make()
        self.entries[0]["version"] = "changed"
        (self.root / "catalog.json").write_text(json.dumps(self.entries))
        with self.assertRaisesRegex(ValueError, "changed"):
            retriever.search("brakes")

    def test_path_traversal_rejected(self):
        self.add("brakes.md", "Brake air leaks")
        self.entries[0]["source"] = "../outside.txt"
        (self.root / "catalog.json").write_text(json.dumps(self.entries))
        with self.assertRaisesRegex(ValueError, "inside"):
            self.make()

    def test_unreviewed_document_rejected(self):
        self.add("brakes.md", "Brake air leaks")
        self.entries[0]["reviewed"] = False
        (self.root / "catalog.json").write_text(json.dumps(self.entries))
        with self.assertRaisesRegex(ValueError, "reviewed"):
            self.make()

    def test_token_chunks_cover_end_of_long_document(self):
        words = ["brake" + str(i) for i in range(100)]
        self.add("long.txt", " ".join(words))
        retriever = self.make()
        self.assertTrue(any("brake99" in c["text"] for c in retriever.chunks))
        self.assertTrue(all(len(c["text"].split()) <= 30 for c in retriever.chunks))

    def test_stable_citation_ids(self):
        self.add("brakes.md", "Brake air leaks")
        first = [h["citation_id"] for h in self.make().chunks]
        second = [h["citation_id"] for h in self.make().chunks]
        self.assertEqual(first, second)

    def test_chinese_tokenization(self):
        self.assertIn("制动", tokens("检查制动系统 Risk6"))
        self.assertIn("risk6", tokens("检查制动系统 Risk6"))

    def test_pdf_page_provenance(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(200, 200)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 100 Td (Brake inspection) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        with (self.root / 'test.pdf').open('wb') as out:
            writer.write(out)
        units = list(document_units(self.root / 'test.pdf'))
        self.assertEqual(units[0][0]['page'], 1)
        self.assertIn('Brake inspection', units[0][1])

    def test_docx_table_ingestion(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Brake inspection')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Leak'
        table.cell(0, 1).text = 'Report'
        doc.save(self.root / 'test.docx')
        units = list(document_units(self.root / 'test.docx'))
        self.assertIn('Leak | Report', units[1][1])
        self.assertEqual(units[1][0]['body_block'], 2)


HIT = {"citation_id": "K-123456abcdef", "kind": "maintenance", "source": "brakes.md",
       "title": "Brakes", "text": "Check for air leaks", "section": "Brakes"}


class Retriever:
    def search(self, query, **kwargs):
        return {"success": True, "hits": [dict(HIT)]}


class Client:
    def __init__(self, replies):
        self.replies, self.requests = iter(replies), []
        self.responses = self

    def create(self, **kwargs):
        self.requests.append(kwargs)
        return next(self.replies)


def response(*calls, answer=""):
    return SimpleNamespace(id="response-id", output=list(calls), output_text=answer)


def call(name, arguments, identifier="call-id"):
    return SimpleNamespace(type="function_call", name=name, arguments=json.dumps(arguments), call_id=identifier)


def vehicle(identifier):
    return {"success": True, "vehicle_id": identifier,
            "current": {"analysis_id": 42, "state": "C4", "action": "PRIORITY_INSPECT", "rule_id": "TR1"}}


class AgentTests(unittest.TestCase):
    def make(self, replies, **kwargs):
        client = Client(replies)
        assistant = FleetCareAssistant(client=client, retriever=Retriever(), vehicle_lookup=vehicle, **kwargs)
        return assistant, client

    def test_two_different_tool_rounds_complete(self):
        assistant, client = self.make([
            response(call("get_vehicle_analysis", {"vehicle_id": 1492})),
            response(call("search_maintenance_knowledge", {"query": "brake inspection"})),
            response(answer="C4 [DB:42]. Check [K-123456abcdef].")])
        result = assistant.ask("Explain vehicle 1492")
        self.assertEqual(len(client.requests), 3)
        output = json.loads(client.requests[1]["input"][0]["output"])
        self.assertEqual(output["vehicle_analysis"]["current"]["action"], "PRIORITY_INSPECT")
        self.assertIn("knowledge", output)
        self.assertTrue(result["sources"][0]["cited_in_answer"])

    def test_explicit_vehicle_prefetch(self):
        assistant, client = self.make([response(answer="C4")])
        result = assistant.ask("Why inspect?", vehicle_id=1492)
        initial = json.loads(client.requests[0]["input"])
        self.assertEqual(initial["requested_vehicle"]["vehicle_analysis"]["vehicle_id"], 1492)
        self.assertEqual(len(result["vehicle_analysis"]), 1)

    def test_multiple_calls_all_get_results(self):
        assistant, client = self.make([
            response(call("get_vehicle_analysis", {"vehicle_id": 1}, "a"),
                     call("get_vehicle_analysis", {"vehicle_id": 2}, "b")), response(answer="Done")])
        assistant.ask("Explain vehicles 1 and 2")
        self.assertEqual([o["call_id"] for o in client.requests[1]["input"]], ["a", "b"])

    def test_tool_limit_forces_final_text(self):
        assistant, client = self.make([
            response(call("get_vehicle_analysis", {"vehicle_id": 1})), response(answer="Done")], max_tool_rounds=1)
        assistant.ask("Explain vehicle 1")
        self.assertEqual(client.requests[-1]["tool_choice"], "none")

    def test_unknown_tool_gets_error_output(self):
        assistant, client = self.make([response(call("delete_vehicle", {"vehicle_id": 1})), response(answer="Unavailable")])
        assistant.ask("Explain vehicle 1")
        result = json.loads(client.requests[1]["input"][0]["output"])
        self.assertFalse(result["success"])

    def test_missing_knowledge_preserves_vehicle_result(self):
        assistant, client = self.make([response(answer="C4")])
        def fail(*args, **kwargs):
            raise FileNotFoundError("index missing")
        assistant.retriever = SimpleNamespace(search=fail)
        result = assistant.ask("Explain", vehicle_id=1)
        self.assertTrue(result["vehicle_analysis"][0]["success"])
        self.assertTrue(result["retrieval_warnings"])

    def test_citation_ids_validated(self):
        assistant, client = self.make([response(answer="Wrong [K-000000000000] [DB:999].")])
        result = assistant.ask("Explain", vehicle_id=1)
        self.assertNotIn("[K-000000000000]", result["answer"])
        self.assertNotIn("[DB:999]", result["answer"])
        self.assertEqual(len(result["retrieval_warnings"]), 2)

    def test_request_evidence_does_not_leak(self):
        assistant, client = self.make([response(answer="One"), response(answer="Two")])
        assistant.ask("Explain", vehicle_id=1)
        result = assistant.ask("General question")
        self.assertEqual(result["vehicle_analysis"], [])

    def test_missing_vehicle_keeps_failure(self):
        assistant, client = self.make([response(answer="Missing")])
        assistant.vehicle_lookup = lambda identifier: {"success": False, "vehicle_id": identifier, "error": "not found"}
        result = assistant.ask("Explain", vehicle_id=999)
        self.assertFalse(result["vehicle_analysis"][0]["success"])

    def test_string_vehicle_id_rejected(self):
        assistant, client = self.make([response(answer="None")])
        with self.assertRaisesRegex(ValueError, "integer"):
            assistant.ask("Explain", vehicle_id="1")


if __name__ == "__main__":
    unittest.main()
